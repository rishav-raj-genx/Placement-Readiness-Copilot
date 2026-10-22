import json

import pymupdf
from fastapi.testclient import TestClient
from pytest import MonkeyPatch

from ai_reviewer import get_provider_sequence
from main import app

client = TestClient(app)


def create_two_page_resume() -> bytes:
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "EXPERIENCE\n"
        "Built and launched a student placement dashboard used by 200 candidates.\n"
        "Led development of typed APIs and reduced response time by 35 percent.\n"
        "SKILLS\nPython FastAPI PostgreSQL Git GitHub Docker testing automation.",
    )
    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "PROJECTS\n"
        "Created a profile review service and improved parsing reliability.\n"
        "EDUCATION\nBachelor degree with software engineering experiance.\n"
        "Successfully delivered teh project with documented results and measurable impact.",
    )
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def test_review_resume_flags_two_page_pdf(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "static")
    monkeypatch.setenv("INTERNAL_SERVICE_KEY", "test-secret-key-1234")
    response = client.post(
        "/review/resume",
        headers={"Authorization": "Bearer test-secret-key-1234"},
        files={
            "file": (
                "two-page-resume.pdf",
                create_two_page_resume(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["page_count"] == 2
    assert payload["is_one_page"] is False
    assert payload["uses_action_verbs"] is True
    assert payload["has_spelling_or_grammar_issues"] is True
    assert "Resume is 2 pages instead of 1." in payload["specific_fixes"]
    print(json.dumps(payload, indent=2))


def test_review_linkedin_pdf_without_resume(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "static")
    monkeypatch.setenv("INTERNAL_SERVICE_KEY", "test-secret-key-1234")
    response = client.post(
        "/review/linkedin",
        headers={"Authorization": "Bearer test-secret-key-1234"},
        files={"file": ("linkedin.pdf", create_two_page_resume(), "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert isinstance(payload["score"], int)
    assert isinstance(payload["specific_fixes"], list)


def test_forced_profile_refresh_isolates_invalid_linkedin_pdf(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "static")
    monkeypatch.setenv("INTERNAL_SERVICE_KEY", "test-secret-key-1234")
    response = client.post(
        "/analyze-profile",
        headers={"Authorization": "Bearer test-secret-key-1234"},
        data={
            "github_username": "",
            "leetcode_url": "",
            "force_refresh": "true",
        },
        files={
            "resume": ("resume.pdf", create_two_page_resume(), "application/pdf"),
            "linkedin_pdf": ("linkedin.txt", b"not-a-pdf", "text/plain"),
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["resume_review"]["score"] >= 0
    assert payload["linkedin_review"] is None
    assert payload["linkedin_error"] == "Upload a PDF or DOCX file with a supported content type."


def test_auth_fails_closed_when_secret_missing(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.delenv("INTERNAL_SERVICE_KEY", raising=False)
    monkeypatch.delenv("INTERNAL_SERVICE_SECRET", raising=False)
    monkeypatch.delenv("PROFILE_REVIEW_SERVICE_TOKEN", raising=False)
    monkeypatch.setenv("E2E_MODE", "false")
    response = client.post(
        "/review/resume",
        files={"file": ("resume.pdf", create_two_page_resume(), "application/pdf")},
    )
    assert response.status_code == 500


def test_auth_rejects_unauthorized_token(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("INTERNAL_SERVICE_SECRET", "test-secret-key-1234")
    monkeypatch.setenv("E2E_MODE", "false")
    response = client.post(
        "/review/resume",
        headers={"Authorization": "Bearer wrong-token"},
        files={"file": ("resume.pdf", create_two_page_resume(), "application/pdf")},
    )
    assert response.status_code == 403


def test_auth_accepts_internal_service_secret_with_whitespace(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("INTERNAL_SERVICE_SECRET", "  my-secret-key-xyz  \n")
    monkeypatch.setenv("INTERNAL_SERVICE_KEY", "fallback-key")
    monkeypatch.setenv("E2E_MODE", "false")
    response = client.post(
        "/review/resume",
        headers={"Authorization": "Bearer  my-secret-key-xyz \n"},
        files={"file": ("resume.pdf", create_two_page_resume(), "application/pdf")},
    )
    assert response.status_code == 200


def test_provider_toggle_and_ollama_fallback(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.delenv("USE_LOCAL_MODE", raising=False)
    monkeypatch.setenv("OLLAMA_FALLBACK", "true")

    monkeypatch.setenv("LLM_PROVIDER", "groq")
    assert get_provider_sequence() == ["groq", "ollama"]

    monkeypatch.setenv("LLM_PROVIDER", "openai")
    assert get_provider_sequence() == ["openai", "ollama"]

    monkeypatch.setenv("LLM_PROVIDER", "ollama")
    assert get_provider_sequence() == ["ollama"]


def create_docx_resume() -> bytes:
    import io
    import docx

    doc = docx.Document()
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Engineered distributed streaming pipeline handling 50k events/sec, reducing latency by 45%.")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, React, TypeScript, Next.js, PostgreSQL, Docker, AWS, Git.")
    doc.add_paragraph("PROJECTS")
    doc.add_paragraph("Built GitHub portfolio at github.com/testuser/project with 99.9% uptime.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_review_resume_docx(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "static")
    monkeypatch.setenv("INTERNAL_SERVICE_KEY", "test-secret-key-1234")
    docx_bytes = create_docx_resume()
    response = client.post(
        "/review/resume",
        headers={"Authorization": "Bearer test-secret-key-1234"},
        files={
            "file": (
                "resume.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["page_count"] == 1
    assert payload["is_one_page"] is True
    assert payload["uses_action_verbs"] is True
    assert payload["has_quantifiable_metrics"] is True


def test_forced_refresh_isolates_corrupted_resume_and_service_failures(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setenv("LLM_PROVIDER", "static")
    monkeypatch.setenv("INTERNAL_SERVICE_KEY", "test-secret-key-1234")
    response = client.post(
        "/analyze-profile",
        headers={"Authorization": "Bearer test-secret-key-1234"},
        data={
            "github_username": "nonexistent_github_user_404_abc",
            "leetcode_url": "https://leetcode.com/nonexistent_leetcode_user_404_abc",
            "force_refresh": "true",
        },
        files={
            "resume": ("corrupted.pdf", b"corrupted-pdf-bytes", "application/pdf"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["resume_review"]["score"] >= 0
    assert payload["resume_error"] is not None
    assert payload["total_readiness_score"] >= 0

