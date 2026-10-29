import io
import os
import re
from collections.abc import Mapping
from typing import Literal, Optional

import docx
from dotenv import load_dotenv
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.prompts import ChatPromptTemplate
from langchain_groq import ChatGroq
from langchain_ollama import ChatOllama
from langchain_openai import ChatOpenAI
from pydantic import BaseModel, Field

load_dotenv()


def extract_docx_data(file_bytes: bytes) -> tuple[str, dict[str, object]]:
    """Extract text and layout geometry metadata from a DOCX byte payload."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    table_texts: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                table_texts.append(" | ".join(cells))

    full_text = "\n".join(paragraphs + table_texts)
    table_count = len(doc.tables)
    margins: dict[str, float] = {}
    has_narrow_margins = False
    if doc.sections:
        section = doc.sections[0]
        left_pt = getattr(getattr(section, "left_margin", None), "pt", 72.0)
        right_pt = getattr(getattr(section, "right_margin", None), "pt", 72.0)
        top_pt = getattr(getattr(section, "top_margin", None), "pt", 72.0)
        bottom_pt = getattr(getattr(section, "bottom_margin", None), "pt", 72.0)
        margins = {
            "left": round(float(left_pt), 2),
            "right": round(float(right_pt), 2),
            "top": round(float(top_pt), 2),
            "bottom": round(float(bottom_pt), 2),
        }
        if float(left_pt) < 18.0 or float(right_pt) < 18.0:
            has_narrow_margins = True

    layout_geometry: dict[str, object] = {
        "format": "docx",
        "table_count": table_count,
        "has_tables": table_count > 0,
        "margins": margins,
        "has_narrow_margins": has_narrow_margins,
        "is_multi_column": table_count >= 2,
    }
    return full_text, layout_geometry


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract all text paragraphs from a DOCX byte payload."""
    text, _ = extract_docx_data(file_bytes)
    return text

ModelProvider = Literal["groq", "openai", "ollama", "static"]


class ReviewResult(BaseModel):
    score: int = Field(ge=0, le=100)
    page_count: int = Field(ge=0)
    is_one_page: bool
    is_ats_friendly: bool
    uses_action_verbs: bool
    has_spelling_or_grammar_issues: bool
    has_coding_profile_links: bool = Field(
        default=False, description="Whether GitHub or LeetCode/coding profile links are included"
    )
    has_project_stack_and_impact: bool = Field(
        default=False, description="Whether projects detail tech stack, impact metrics, and repository/live links"
    )
    has_clean_visual_layout: bool = Field(
        default=True, description="Whether visual layout and structure are clean, readable, and free of clutter"
    )
    has_quantifiable_metrics: bool = Field(
        default=False, description="Whether bullet points contain quantifiable metrics (e.g., percentages, latencies, user counts)"
    )
    has_ats_keywords: bool = Field(
        default=False, description="Whether resume contains industry-standard ATS technical keywords"
    )
    spelling_grammar_issues: list[str] = Field(default_factory=list)
    specific_fixes: list[str] = Field(default_factory=list)


class GithubReview(BaseModel):
    score: int = Field(ge=0, le=100)
    has_pinned_repos: bool = Field(default=False, description="Whether representative repositories are pinned")
    has_recent_commit_activity: bool = Field(default=False, description="Whether active commits occurred in last 90 days")
    has_committed_secrets: bool = Field(default=False, description="Whether committed API keys or env files were found")
    has_unit_tests: bool = Field(default=False, description="Whether repositories include unit or integration test suites")
    has_ci_cd: bool = Field(default=False, description="Whether repositories include CI/CD pipelines (e.g. GitHub Actions workflows)")
    has_architecture_diagrams: bool = Field(default=False, description="Whether READMEs contain architecture or system workflow diagrams")
    code_quality_score: int = Field(default=0, ge=0, le=100, description="Score based on structure, linters, CI, and tests")
    specific_fixes: list[str] = Field(default_factory=list)


class LinkedInReviewResult(BaseModel):
    score: int = Field(ge=0, le=100)
    has_complete_profile: bool
    skills_relevant: bool
    has_photo: bool = Field(default=False, description="Whether professional profile photo / headshot is detected")
    has_custom_url: bool = Field(default=False, description="Whether a custom vanity LinkedIn URL is configured")
    has_engagement_markers: bool = Field(default=False, description="Whether profile exhibits network engagement (posts, recommendations, activity)")
    is_copy_paste_of_resume: bool = Field(default=False, description="Whether profile is merely a verbatim copy-paste of resume without network context")
    aligned_with_resume: bool = Field(default=True, description="Whether experiences and target roles align with resume")
    suggested_rewrites: list[str] = Field(
        default_factory=list, description="Ready-to-paste rewrites for headline, about section, and key bullet points"
    )
    specific_fixes: list[str] = Field(
        default_factory=list, description="Specific actionable fixes to improve LinkedIn profile"
    )


class ReviewBackendError(RuntimeError):
    """Raised when every configured LLM backend fails."""


def _enabled(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


def get_provider_sequence() -> list[ModelProvider]:
    configured = os.getenv("LLM_PROVIDER", "auto").strip().lower()
    if _enabled("USE_LOCAL_MODE"):
        configured = "ollama"

    valid_providers = {"auto", "groq", "openai", "ollama", "static"}
    if configured not in valid_providers:
        raise ValueError(
            "LLM_PROVIDER must be one of: auto, groq, openai, ollama, static."
        )

    if configured == "static":
        return ["static"]
    if configured == "ollama":
        return ["ollama"]

    providers: list[ModelProvider] = []
    if configured == "auto":
        preferred_api = os.getenv("API_LLM_PROVIDER", "groq").strip().lower()
        if preferred_api == "openai" and os.getenv("OPENAI_API_KEY"):
            providers.append("openai")
        elif os.getenv("GROQ_API_KEY"):
            providers.append("groq")
        elif os.getenv("OPENAI_API_KEY"):
            providers.append("openai")
    elif configured == "groq":
        providers.append("groq")
    elif configured == "openai":
        providers.append("openai")

    if _enabled("OLLAMA_FALLBACK", default=True):
        providers.append("ollama")

    return providers or ["ollama"]


def get_backend_status() -> dict[str, object]:
    return {
        "configured_provider": os.getenv("LLM_PROVIDER", "auto").strip().lower(),
        "provider_sequence": get_provider_sequence(),
        "ollama_fallback_enabled": _enabled("OLLAMA_FALLBACK", default=True),
    }


def _get_chat_model(provider: ModelProvider) -> BaseChatModel:
    temperature = 0
    if provider == "groq":
        return ChatGroq(
            model=os.getenv("GROQ_MODEL", "llama-3.1-8b-instant"),
            temperature=temperature,
        )
    if provider == "openai":
        return ChatOpenAI(
            model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
            temperature=temperature,
        )
    if provider == "ollama":
        return ChatOllama(
            model=os.getenv("OLLAMA_MODEL", "llama3.1:8b"),
            base_url=os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434"),
            temperature=temperature,
        )
    raise ValueError("The static provider does not create a chat model.")


def _sanitize_error_message(msg: str) -> str:
    """Sanitize upstream provider errors to avoid leaking credentials or sensitive payloads."""
    sanitized = re.sub(
        r"(?i)(bearer|key|token|auth|password|secret)[\s:=]+['\"]?[a-zA-Z0-9_\-\.]{8,}['\"]?",
        r"\1=[REDACTED]",
        msg,
    )
    sanitized = re.sub(r"https?://[^\s\?]+(\?[^\s]*)?", "[URL]", sanitized)
    if len(sanitized) > 100:
        sanitized = sanitized[:100] + "..."
    return sanitized


def _invoke_structured[SchemaType: BaseModel](
    prompt: ChatPromptTemplate,
    schema: type[SchemaType],
    payload: Mapping[str, object],
) -> SchemaType:
    failures: list[str] = []
    for provider in get_provider_sequence():
        if provider == "static":
            continue
        try:
            reviewer = prompt | _get_chat_model(provider).with_structured_output(schema)
            response = reviewer.invoke(dict(payload))
            return response if isinstance(response, schema) else schema.model_validate(response)
        except Exception as exc:  # noqa: BLE001
            failures.append(f"{provider}: {_sanitize_error_message(str(exc))}")

    raise ReviewBackendError("; ".join(failures) or "No LLM provider is enabled.")


ACTION_VERBS = {
    "achieved",
    "automated",
    "built",
    "created",
    "delivered",
    "designed",
    "developed",
    "implemented",
    "improved",
    "increased",
    "launched",
    "led",
    "managed",
    "optimized",
    "reduced",
    "resolved",
    "scaled",
    "shipped",
}

COMMON_MISSPELLINGS = {
    "algoritm": "algorithm",
    "algoritms": "algorithms",
    "architechture": "architecture",
    "commited": "committed",
    "compatability": "compatibility",
    "databse": "database",
    "databses": "databases",
    "dependancy": "dependency",
    "developement": "development",
    "enviroment": "environment",
    "excede": "exceed",
    "experiance": "experience",
    "implimented": "implemented",
    "independant": "independent",
    "infrastucture": "infrastructure",
    "maintainance": "maintenance",
    "managment": "management",
    "occured": "occurred",
    "occurrance": "occurrence",
    "performence": "performance",
    "persistant": "persistent",
    "priviledge": "privilege",
    "proccess": "process",
    "recieved": "received",
    "refering": "referring",
    "reponsible": "responsible",
    "responsibile": "responsible",
    "seperate": "separate",
    "succesful": "successful",
    "succesfully": "successfully",
    "teh": "the",
}


def _check_grammar_and_style(text: str) -> list[str]:
    """Perform comprehensive grammatical, stylistic, and technical typo verification."""
    issues: list[str] = []
    normalized = text.lower()
    words = re.findall(r"[a-zA-Z']+", normalized)

    # 1. Spelling verification against technical and professional vocabulary
    for misspelling, correction in COMMON_MISSPELLINGS.items():
        if misspelling in words:
            issues.append(f"Replace '{misspelling}' with '{correction}'.")

    # 2. Repeated words / stuttering
    repeated = re.findall(
        r"\b(the|and|in|with|to|for|is|that|on|at|of|a|an)\s+\1\b", normalized
    )
    for rep in sorted(set(repeated)):
        issues.append(f"Grammar: Repeated word '{rep} {rep}' detected.")

    # 3. Preposition, collocation, and tense issues
    if re.search(r"\bresponsible\s+of\b", normalized):
        issues.append("Grammar: Replace 'responsible of' with 'responsible for'.")
    if re.search(r"\bexperiences\s+in\b", normalized):
        issues.append("Grammar: Use uncountable 'experience in' rather than 'experiences in'.")
    if re.search(r"\b(?:did\s+not|to|will)\s+loose\b|\bloose\s+(?:data|users|performance|access|state)\b", normalized):
        issues.append("Word choice: Replace 'loose' with 'lose'.")
    if re.search(r"\beffected\s+(?:users|performance|latency|throughput|systems?)\b", normalized):
        issues.append("Word choice: Use 'affected' (verb) instead of 'effected'.")
    if re.search(r"\bthere\s+(?:work|code|systems?|projects?|team)\b", normalized):
        issues.append("Grammar: Check 'there' vs possessive 'their'.")

    # 4. Spacing after commas
    punctuation_spacing = re.findall(r"\b[a-zA-Z]+,[a-zA-Z]+\b", text)
    if punctuation_spacing:
        issues.append(f"Punctuation: Missing space after comma (e.g., '{punctuation_spacing[0]}').")

    # 5. Bullet point capitalization
    for line in text.splitlines():
        trimmed = line.strip()
        if trimmed.startswith(("-", "*", "•")):
            bullet_body = trimmed.lstrip("-*• ").strip()
            if bullet_body and bullet_body[0].islower():
                issues.append(f"Style: Capitalize first word of bullet point: '{bullet_body[:30]}...'")
                break

    return issues


ATS_CORE_KEYWORDS = {
    "react", "next.js", "python", "typescript", "javascript", "node", "docker",
    "kubernetes", "postgres", "sql", "redis", "aws", "gcp", "azure", "git",
    "ci/cd", "rest", "graphql", "system design", "algorithms", "microservices"
}


def _static_resume_review(
    text: str, page_count: int, layout_geometry: Optional[dict[str, object]] = None
) -> ReviewResult:
    normalized = text.lower()
    words = re.findall(r"[a-zA-Z']+", normalized)
    section_matches = sum(
        heading in normalized
        for heading in ("experience", "education", "skills", "projects")
    )
    unusual_character_count = sum(
        not (character.isprintable() or character in "\n\t") for character in text
    )

    is_one_page = page_count == 1
    is_ats_friendly = len(words) >= 40 and section_matches >= 2 and unusual_character_count == 0
    uses_action_verbs = any(word in ACTION_VERBS for word in words)

    # 1. Coding Profile Links Check
    has_links = bool(re.search(r"(github\.com|leetcode\.com|gitlab\.com|portfolio)", normalized))

    # 2. Quantifiable Metrics Check (e.g., "Increased speed by 20%", "500 req/sec", "$12K")
    has_metrics = bool(re.search(
        r"\b\d+(?:\.\d+)?%|\b\d+x\b|\b\d+[\s]*(?:users|req(?:uests)?|rps|qps|ms|seconds|minutes|k|m|million|billion|clients|sessions|tps)\b|\$\d+",
        normalized
    ))

    # 3. Project Stack, Impact, and Links Check
    has_tech_keywords = any(kw in normalized for kw in ("react", "next.js", "python", "typescript", "postgres", "node", "docker", "api"))
    has_project_stack = "project" in normalized and has_tech_keywords and has_metrics

    # 4. ATS Technical Keywords Check
    ats_matches = sum(kw in normalized for kw in ATS_CORE_KEYWORDS)
    has_ats_keywords = ats_matches >= 3

    # 5. Generic Summary Check
    has_generic_summary = bool(re.search(
        r"\b(seeking an entry[- ]level|motivated student|hardworking individual|quick learner|eager to learn|enthusiastic developer)\b",
        normalized
    ))

    # 6. Clean Visual Layout & Geometry Check
    has_narrow_margins = bool(layout_geometry.get("has_narrow_margins")) if layout_geometry else False
    is_multi_column = bool(layout_geometry.get("is_multi_column")) if layout_geometry else False

    has_clean_layout = unusual_character_count == 0 and section_matches >= 2 and not has_narrow_margins
    if is_multi_column:
        is_ats_friendly = False

    # 7. Comprehensive Grammar and Spelling
    grammar_and_spelling_issues = _check_grammar_and_style(text)
    has_language_issues = bool(grammar_and_spelling_issues)

    fixes: list[str] = [
        "[Heuristic Check] Rule-based evaluation used. Visual layout, columns, tables, and complex grammar verified structurally."
    ]

    # Ruthless Cap: If bullet points lack quantifiable metrics, cap score at 60/100
    if not has_metrics:
        fixes.insert(0, "[Score Capped at 60/100] Bullet points lack quantifiable metrics. Top-tier recruiters reject vague statements like 'improved speed'. Quantify engineering outcomes with numbers, percentages, or latencies (e.g., 'Reduced p99 latency by 35% across 10,000 active sessions').")

    if not is_one_page:
        fixes.append(f"Resume is {page_count} pages instead of 1.")
    if not is_ats_friendly:
        fixes.append("Use standard section headings, selectable text, and a single-column ATS-friendly layout.")
    if is_multi_column:
        fixes.append("Layout Geometry: Multi-column structure detected. ATS parsers read horizontally across columns, corrupting experience timelines. Convert to a single-column layout.")
    if has_narrow_margins:
        fixes.append("Layout Geometry: Document margins are too narrow (<0.25in / 18pt), risking text clipping and parser rejection. Maintain standard 0.5in–1.0in margins.")
    if not uses_action_verbs:
        fixes.append("Start achievement bullets with strong action verbs and quantify outcomes.")
    if not has_links:
        fixes.append("Add prominent clickable GitHub and LeetCode/coding profile URLs to your contact header.")
    if not has_project_stack:
        fixes.append("Detail technologies used, measurable performance impact, and repository/live demo links for each project.")
    if has_generic_summary:
        fixes.append("Deducted points for generic summary. Remove clichés like 'hardworking student' or 'seeking an entry-level role' and replace with a targeted technical skills profile.")
    if not has_ats_keywords:
        fixes.append("Missing core ATS keywords for software engineering. Ensure industry-standard skills (languages, databases, frameworks, CI/CD) are explicitly represented in a Skills section.")
    if not has_clean_layout and not has_narrow_margins:
        fixes.append("Improve visual hierarchy with clear section dividers, consistent font sizes, and uncluttered margins.")
    fixes.extend(grammar_and_spelling_issues)

    # Balanced score out of 100
    points = 0
    if is_one_page:
        points += 20
    if is_ats_friendly:
        points += 20
    if uses_action_verbs:
        points += 15
    if not has_language_issues:
        points += 15
    if has_links:
        points += 10
    if has_project_stack:
        points += 10
    if has_ats_keywords:
        points += 10

    if has_generic_summary:
        points = max(0, points - 15)

    final_score = min(100, points)
    # Strictly enforce 60/100 cap if bullet points lack quantifiable metrics
    if not has_metrics:
        final_score = min(final_score, 60)

    return ReviewResult(
        score=final_score,
        page_count=page_count,
        is_one_page=is_one_page,
        is_ats_friendly=is_ats_friendly,
        uses_action_verbs=uses_action_verbs,
        has_spelling_or_grammar_issues=has_language_issues,
        has_coding_profile_links=has_links,
        has_project_stack_and_impact=has_project_stack,
        has_clean_visual_layout=has_clean_layout,
        has_quantifiable_metrics=has_metrics,
        has_ats_keywords=has_ats_keywords,
        spelling_grammar_issues=grammar_and_spelling_issues,
        specific_fixes=fixes,
    )


def evaluate_text(
    text: str, page_count: int, layout_geometry: Optional[dict[str, object]] = None
) -> ReviewResult:
    static_review = _static_resume_review(text, page_count, layout_geometry=layout_geometry)
    if get_provider_sequence() == ["static"]:
        return static_review

    prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                """You are a Principal Technical Recruiter and FAANG Hiring Manager conducting a ruthless ATS and resume audit.
Evaluate the candidate's resume strictly against top-tier engineering standards:
1. QUANTIFIABLE METRICS (MANDATORY CAP AT 60):
   - Bullet points MUST contain quantifiable metrics (e.g., "Increased speed by 20%", "Decreased p99 latency by 35%", "Scaled throughput to 10k RPS").
   - If bullet points lack quantifiable metrics or rely on generic statements (e.g. "Made it faster", "Helped team"), YOU MUST CAP THE SCORE AT 60/100 and set has_quantifiable_metrics=False.
2. ATS KEYWORDS & GENERIC SUMMARIES: Deduct 15-20 points for generic introductory summaries or missing ATS technical keywords. Set has_ats_keywords.
3. PAGE COUNT: Must be exactly one page for early career/students. Cap at 75 if > 1 page.
4. CODING PROFILE LINKS: Explicitly check for verified GitHub, LeetCode, or portfolio URLs. Set has_coding_profile_links.
5. PROJECT STACK & IMPACT: Flag basic CRUD apps without clear architecture, tech stack depth, or deployed links. Set has_project_stack_and_impact.
6. VISUAL LAYOUT & ATS COMPLIANCE: Single-column, no multi-column tables/graphics that break ATS parsers. Set has_clean_visual_layout.
7. SPELLING & GRAMMAR: Zero tolerance for typos in technical terms or grammar slips.

Provide actionable, high-standard recommendations in specific_fixes. Return only the ReviewResult schema.""",
            ),
            (
                "human",
                "Verified PDF page_count: {page_count}\nLayout geometry: {layout_geometry}\n\nExtracted resume text:\n{resume_text}",
            ),
        ]
    )

    try:
        review = _invoke_structured(
            prompt,
            ReviewResult,
            {
                "page_count": page_count,
                "resume_text": text,
                "layout_geometry": str(layout_geometry or {}),
            },
        )
    except ReviewBackendError:
        return static_review

    review.page_count = page_count
    review.is_one_page = page_count == 1
    page_fix = f"Resume is {page_count} pages instead of 1."
    if not review.is_one_page and page_fix not in review.specific_fixes:
        review.specific_fixes.insert(0, page_fix)
    if not review.is_one_page:
        review.score = min(review.score, 75)

    if layout_geometry:
        if layout_geometry.get("has_narrow_margins"):
            review.has_clean_visual_layout = False
            margin_fix = "Layout Geometry: Document margins are too narrow (<0.25in / 18pt), risking text clipping and parser rejection. Maintain standard 0.5in–1.0in margins."
            if margin_fix not in review.specific_fixes:
                review.specific_fixes.append(margin_fix)
        if layout_geometry.get("is_multi_column"):
            review.is_ats_friendly = False
            col_fix = "Layout Geometry: Multi-column structure detected. ATS parsers read horizontally across columns, corrupting experience timelines. Convert to a single-column layout."
            if col_fix not in review.specific_fixes:
                review.specific_fixes.append(col_fix)

    # Strictly enforce 60/100 cap if bullet points lack quantifiable metrics
    if not review.has_quantifiable_metrics:
        review.score = min(review.score, 60)
        cap_msg = "[Score Capped at 60/100] Bullet points lack quantifiable metrics. Quantify engineering achievements with concrete numbers and percentages."
        if not any("score capped at 60/100" in f.lower() for f in review.specific_fixes):
            review.specific_fixes.insert(0, cap_msg)

    # Merge deterministic grammar and spelling verification
    deterministic_grammar = _check_grammar_and_style(text)
    if deterministic_grammar:
        review.has_spelling_or_grammar_issues = True
        for issue in deterministic_grammar:
            if issue not in review.spelling_grammar_issues:
                review.spelling_grammar_issues.append(issue)
            if issue not in review.specific_fixes:
                review.specific_fixes.append(issue)

    return review


def evaluate_github(github_data: Mapping[str, object]) -> GithubReview:
    prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                """You are a FAANG Principal Software Engineer conducting a ruthless GitHub audit.
Evaluate the candidate's GitHub repositories strictly against top-tier industry standards:
1. MANDATORY 50/100 SCORE CAP:
   - Unit Tests: Repositories must contain automated unit/integration test suites. Set has_unit_tests.
   - CI/CD Pipelines: Repositories must include GitHub Actions workflows (.github/workflows) running automated tests/linting. Set has_ci_cd.
   - Architecture Diagrams: README files must include architecture/system diagrams (Mermaid, plantuml, or visual diagram images) explaining component dataflow. Set has_architecture_diagrams.
   - CAP THE SCORE AT 50/100 if repositories lack Unit Tests, CI/CD pipelines, OR architecture diagrams in the README. A basic CRUD app with a default README is not enough.
2. PINNED REPOSITORIES: Representative flagship projects must be pinned with descriptions and topics. Set has_pinned_repos.
3. RECENT ACTIVITY: Active commit cadence within the last 90 days. Set has_recent_commit_activity.
4. COMMITTED SECRETS: Leaked API keys, database credentials, or .env files must trigger immediate score deduction and critical security warning. Set has_committed_secrets.
5. CODE QUALITY: Clean directory modularity (src/lib), dependency manifests, linters/formatters. Set code_quality_score (0-100).

Return only the GithubReview schema.""",
            ),
            ("human", "{github_data}"),
        ]
    )
    try:
        review = _invoke_structured(
            prompt,
            GithubReview,
            {"github_data": str(dict(github_data))},
        )
        if not (review.has_unit_tests and review.has_ci_cd and review.has_architecture_diagrams):
            review.score = min(review.score, 50)
            if not any("capped at 50" in f.lower() for f in review.specific_fixes):
                review.specific_fixes.insert(0, "[Score Capped at 50/100] Repositories lack Unit Tests, CI/CD pipelines, or README Architecture Diagrams. Top-tier SDE standards require automated test suites, CI/CD, and architecture diagrams.")
        return review
    except ReviewBackendError:
        repositories = github_data.get("repositories", [])
        if not isinstance(repositories, list) or not repositories:
            return GithubReview(
                score=0,
                has_pinned_repos=False,
                has_recent_commit_activity=False,
                has_committed_secrets=False,
                has_unit_tests=False,
                has_ci_cd=False,
                has_architecture_diagrams=False,
                code_quality_score=0,
                specific_fixes=[
                    "[Score Capped at 50/100] No public repositories found. Add publicly visible source code repositories with unit tests, CI/CD pipelines, architecture diagrams, and dependency manifests.",
                ],
            )

        repo_scores: list[float] = []
        fixes: list[str] = [
            "[Heuristic Check] Deterministic repository inspection used (LLM provider offline)."
        ]
        has_tests_any = False
        has_ci_any = False
        has_diagram_any = False
        has_manifest_any = False
        has_readme_any = False
        has_pinned = bool(github_data.get("has_pinned_repos"))
        has_recent = bool(github_data.get("has_recent_commit_activity"))
        has_secrets = bool(github_data.get("has_committed_secrets"))

        for repo in repositories:
            if not isinstance(repo, dict):
                continue
            r_score = 15.0
            structure = repo.get("structure", {}) if isinstance(repo.get("structure"), dict) else {}
            if repo.get("has_readme") and int(repo.get("readme_length", 0)) > 150:
                r_score += 20.0
                has_readme_any = True
            if structure.get("has_source_directory") or structure.get("has_dependency_manifest"):
                r_score += 20.0
                has_manifest_any = True
            if structure.get("has_tests") or repo.get("has_unit_tests"):
                r_score += 15.0
                has_tests_any = True
            if structure.get("has_ci_workflow") or repo.get("has_ci_cd"):
                r_score += 15.0
                has_ci_any = True
            if structure.get("has_architecture_diagram") or repo.get("has_architecture_diagram"):
                r_score += 15.0
                has_diagram_any = True
            if repo.get("description"):
                r_score += 5.0
            if repo.get("is_fork"):
                r_score *= 0.5
            if repo.get("has_committed_secrets"):
                has_secrets = True
                fixes.append(f"CRITICAL: Committed secrets detected in repo '{repo.get('name')}'. Rotate keys and purge from git history immediately.")
            if repo.get("history_secrets"):
                has_secrets = True
                fixes.append(f"CRITICAL: Leaked credentials found in git commit history of '{repo.get('name')}'. Scrub history immediately using git-filter-repo.")
            if repo.get("has_readme") and not repo.get("has_setup_instructions"):
                fixes.append(f"README in '{repo.get('name')}' lacks explicit setup/installation instructions (e.g. 'pnpm install', 'docker compose up').")
            if repo.get("code_readability_comments"):
                for comment in repo.get("code_readability_comments", []):
                    if "Clean" not in comment and "inferred" not in comment:
                        fixes.append(f"Repository '{repo.get('name')}': {comment}")
            repo_scores.append(min(100.0, r_score))

        avg_score = round(sum(repo_scores) / len(repo_scores)) if repo_scores else 0
        if has_secrets:
            avg_score = max(0, avg_score - 40)

        # Cap the score at 50/100 if repos lack Unit Tests, CI/CD pipelines, or architecture diagrams
        if not (has_tests_any and has_ci_any and has_diagram_any):
            avg_score = min(avg_score, 50)
            if not has_tests_any:
                fixes.insert(0, "[Score Capped at 50/100] Repositories lack automated Unit Tests. A basic CRUD app without tests will not pass top-tier SDE screening. Add unit/integration tests with Jest, Pytest, or Vitest.")
            if not has_ci_any:
                fixes.insert(0, "[Score Capped at 50/100] Missing CI/CD pipelines. Repositories must have GitHub Actions workflows (.github/workflows) running tests and linters automatically on push.")
            if not has_diagram_any:
                fixes.insert(0, "[Score Capped at 50/100] README lacks architecture diagrams. Include system architecture diagrams (Mermaid or visual flowcharts) illustrating system design and data flow.")

        if not has_pinned:
            fixes.append("Pin your top 2-3 flagship repositories to your GitHub profile overview.")
        if not has_recent:
            fixes.append("No commits detected in the last 90 days. Maintain regular commit activity on your projects.")
        if not has_readme_any:
            fixes.append("Add detailed README files (>150 characters) explaining architecture, setup, and demo links.")
        if not has_manifest_any:
            fixes.append("Include standard dependency manifests (package.json, requirements.txt, or pyproject.toml).")

        quality_score = int(github_data.get("average_code_quality_score", 50))

        return GithubReview(
            score=avg_score,
            has_pinned_repos=has_pinned,
            has_recent_commit_activity=has_recent,
            has_committed_secrets=has_secrets,
            has_unit_tests=has_tests_any,
            has_ci_cd=has_ci_any,
            has_architecture_diagrams=has_diagram_any,
            code_quality_score=quality_score,
            specific_fixes=fixes,
        )


def evaluate_linkedin_text(text: str, resume_context: Optional[str] = None) -> LinkedInReviewResult:
    normalized = text.lower()
    has_photo_hint = bool(re.search(r"\b(photo|picture|headshot|avatar|image|profile photo)\b", normalized))
    has_custom_url = bool(re.search(r"linkedin\.com/in/[a-zA-Z0-9_-]{3,}", normalized)) and not bool(re.search(r"-[0-9a-f]{8,}\b", normalized))
    has_engagement = bool(re.search(r"\b(recommendation|endorsed|post|article|activity|featured|network|mentoring|volunteer|community)\b", normalized))

    prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                """You are an executive tech recruiter and LinkedIn optimization expert for top-tier tech companies.
Evaluate the candidate's LinkedIn profile ruthlessly against industry standards:
1. CUSTOM VANITY URL (MANDATORY): Must use a customized LinkedIn URL (e.g., linkedin.com/in/alex-morgan instead of default random numbers linkedin.com/in/alex-12849182/). Set has_custom_url.
2. KEYWORD-RICH HEADLINE: The headline must include specific target roles and core technical keywords (e.g. "Full-Stack Engineer | TypeScript, React, Distributed Systems"), not generic student titles ("Student at XYZ University").
3. DETAILED EXPERIENCE: Role sections must detail concrete engineering achievements with quantified impact, not just job titles.
4. RESUME COPY-PASTE & NETWORK ENGAGEMENT (MANDATORY CAP AT 70):
   - If the profile is merely a copy-paste of the resume without network engagement markers (such as peer recommendations, featured articles/posts, open-source project showcases, or community activity), YOU MUST CAP THE SCORE AT 70/100. Set is_copy_paste_of_resume=True and has_engagement_markers=False.
5. READY-TO-PASTE REWRITES:
   - Output exact, fully written, ready-to-paste text:
     * "Ready-to-Paste Headline: <complete professional headline with target role and keywords>"
     * "Ready-to-Paste About Section: <complete 1-2 paragraph professional summary ready to copy into LinkedIn>"
     * "Ready-to-Paste Experience Bullet: <exact rewritten bullet with strong action verb and quantified metric>"

Return only the LinkedInReviewResult schema.""",
            ),
            (
                "human",
                "Resume Context (for alignment):\n{resume_context}\n\nLinkedIn Export Text:\n{linkedin_text}",
            ),
        ]
    )

    try:
        review = _invoke_structured(
            prompt,
            LinkedInReviewResult,
            {
                "linkedin_text": text,
                "resume_context": resume_context or "No explicit resume supplied for cross-check.",
            },
        )
        if review.is_copy_paste_of_resume or not review.has_engagement_markers:
            review.score = min(review.score, 70)
            if not any("capped at 70" in f.lower() for f in review.specific_fixes):
                review.specific_fixes.insert(0, "[Score Capped at 70/100] Profile is a direct copy-paste of the resume without network engagement markers. Request peer recommendations and share technical posts to exceed 70.")
        return review
    except ReviewBackendError:
        fixes: list[str] = []
        if len(normalized.strip()) < 120:
            return LinkedInReviewResult(
                score=0,
                has_complete_profile=False,
                skills_relevant=False,
                has_photo=False,
                has_custom_url=False,
                has_engagement_markers=False,
                is_copy_paste_of_resume=False,
                aligned_with_resume=True,
                suggested_rewrites=[
                    "Ready-to-Paste Headline: Software Engineer | React, TypeScript & Node.js | Scalable Systems Builder",
                    "Ready-to-Paste About Section: Motivated Software Engineer with hands-on experience developing full-stack web applications. Passionate about writing clean, maintainable code and building high-performance products.",
                ],
                specific_fixes=[
                    "Profile text is too sparse. Populate headline, detailed experience, and skills sections."
                ],
            )

        has_about = bool(re.search(r"about[:\s\n\r]+.{20,}", normalized, re.DOTALL))
        has_headline = "headline" in normalized or any(
            term in normalized for term in ("engineer", "developer", "student", "intern", "specialist")
        )
        has_skills = "skills" in normalized and bool(re.search(r"skills[:\s\n\r]+.{10,}", normalized, re.DOTALL))
        has_complete = has_about and has_headline

        # Check alignment with resume keywords if provided
        aligned = True
        is_copy_paste = False
        if resume_context:
            res_words = set(re.findall(r"[a-z]{4,}", resume_context.lower()))
            overlap = sum(w in normalized for w in list(res_words)[:30])
            aligned = overlap >= 3
            overlap_ratio = overlap / max(1, len(res_words))
            if (overlap >= 8 or (len(res_words) >= 5 and overlap_ratio >= 0.6)) and not has_engagement:
                is_copy_paste = True

        score = (25 if has_headline else 0) + (25 if has_about else 0) + (20 if has_skills else 0) + (15 if has_photo_hint else 0) + (15 if has_custom_url else 0)

        # Cap at 70/100 if it's just a copy-paste of the resume without network engagement markers
        if is_copy_paste or not has_engagement:
            score = min(score, 70)
            fixes.append("[Score Capped at 70/100] Profile is a copy-paste of the resume without network engagement markers. Add peer recommendations, featured projects, and industry posts to build social proof.")

        if not has_custom_url:
            fixes.append("Configure a clean LinkedIn vanity URL (e.g., linkedin.com/in/yourname) in Public Profile Settings.")

        if not has_photo_hint:
            fixes.append("Upload a professional high-resolution headshot. Profiles with photos receive significantly more recruiter reachouts.")

        rewrites: list[str] = [
            "Ready-to-Paste Headline: Full-Stack Software Engineer | TypeScript, React, Next.js & Python | Building Cloud-Native Products",
            "Ready-to-Paste About Section: Software Engineer dedicated to architecting resilient, user-focused web applications. Proficient in modern frontend ecosystems, RESTful microservices, and relational database design. Eager to solve complex challenges in fast-paced engineering teams.",
            "Ready-to-Paste Experience Bullet: Engineered full-stack features using Next.js and PostgreSQL, decreasing API latency by 35% and supporting over 10,000 active sessions.",
        ]

        return LinkedInReviewResult(
            score=round(score),
            has_complete_profile=has_complete,
            skills_relevant=has_skills,
            has_photo=has_photo_hint,
            has_custom_url=has_custom_url,
            has_engagement_markers=has_engagement,
            is_copy_paste_of_resume=is_copy_paste,
            aligned_with_resume=aligned,
            suggested_rewrites=rewrites,
            specific_fixes=fixes,
        )


def extract_projects_from_resume_text(text: str) -> list[dict[str, object]]:
    """
    Parses and extracts candidate project entries directly from resume text.
    Identifies project titles, descriptions, and repository or demo links.
    """
    if not text or not text.strip():
        return []

    projects: list[dict[str, object]] = []

    pattern = re.compile(
        r"(?:^|\n)\s*(?:(?:academic|personal|key|featured|technical|selected)\s+)?projects?\b[:\s]*(.*?)(?=\n\s*(?:(?:work\s+)?experience|employment|education|skills|certifications|achievements|leadership|\Z))",
        re.DOTALL | re.IGNORECASE,
    )
    match = pattern.search(text)
    section_text = match.group(1).strip() if match else text

    blocks = re.split(r"\n\s*\n|\n(?=[A-Z0-9][A-Za-z0-9\s\-–—|]{3,50}(?:\s*\(|\s*\||\s*–|\s*-|\s*:))", section_text)

    for block in blocks:
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not lines:
            continue
        first_line = lines[0]
        if re.match(r"^(?:projects?|skills|education|experience|summary|objective)$", first_line.lower()):
            lines = lines[1:]
            if not lines:
                continue
            first_line = lines[0]

        title_parts = re.split(r"[|–—\-(:]", first_line)
        title = title_parts[0].strip()
        clean_title = re.sub(r"^[•\-\*\d\.\)]+\s*", "", title).strip()
        if len(clean_title) < 2 or len(clean_title) > 60:
            continue
        if clean_title.lower() in ("project", "projects", "technical skills", "education", "experience", "languages", "tools"):
            continue

        urls = re.findall(r"https?://[^\s)\]]+|github\.com/[a-zA-Z0-9_.-]+/[a-zA-Z0-9_.-]+", block)
        repo_url = None
        live_url = None
        for u in urls:
            clean_u = u if u.startsWith("http") else f"https://{u}"
            if "github.com" in clean_u:
                repo_url = clean_u
            else:
                live_url = clean_u

        desc_lines = lines[1:] if len(lines) > 1 else lines
        desc = " ".join(desc_lines)[:350].strip() or f"Engineering project focused on {clean_title}."

        clean_id = re.sub(r"[^a-z0-9]", "-", clean_title.lower())[:24]
        projects.append({
            "id": f"proj-{clean_id}" if clean_id else f"proj-{len(projects)+1}",
            "name": clean_title,
            "title": clean_title,
            "description": desc,
            "repositoryUrl": repo_url,
            "repoUrl": repo_url,
            "liveUrl": live_url,
        })
        if len(projects) >= 6:
            break

    return projects
